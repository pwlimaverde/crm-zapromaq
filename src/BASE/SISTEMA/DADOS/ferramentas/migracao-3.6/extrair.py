# -*- coding: utf-8 -*-
"""
extrair.py - CRM 3.6 (.xlsm) -> arquivos de carga normalizados (JSON)

Roda na sessao de IA (Linux). Le a planilha por NOME de coluna, nunca por posicao.
Aplica a normalizacao da secao 5.1 do estudo, resolve vinculos, congela os codigos
CT-CCCC-NNNN e AT-CCCC-NNNN, e grava um arquivo por tabela em 40 - DICIONARIO\\carga.

Uso:  python3 extrair.py "<caminho do .xlsm>" "<pasta de saida>"
"""
import sys, os, json, re, unicodedata
from datetime import datetime, date, time
import openpyxl

LINHA_CABECALHO = 3
LINHA_DADOS = 4

# ---------------------------------------------------------------- normalizacao
def trim(v):
    if v is None:
        return ""
    s = str(v).replace(" ", " ").strip()
    return re.sub(r"[ \t]{2,}", " ", s)

def maiusc(v):
    return trim(v).upper()

def minusc(v):
    return trim(v).lower()

def digitos(v):
    return re.sub(r"\D", "", trim(v))

RE_CNPJ_TXT = re.compile(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}")
RE_TEL_TXT  = re.compile(r"\+?\d{0,3}\s*\(?\d{2}\)?[\s.-]?\d{4,5}[\s.-]?\d{4}")

def telefone_partes(v):
    """Devolve (telefone so digitos, nota para observacoes).

    O campo da planilha costuma trazer 'NOME: numero' ou varios numeros.
    O campo do banco guarda SO DIGITOS de um numero; o nome e os numeros
    adicionais vao para observacoes, em vez de virar um numero que nao existe.
    """
    s = str(v or "").replace("\u00a0", " ").strip()
    if not s:
        return "", ""

    limpo = RE_CNPJ_TXT.sub(" ", s)
    achados = []
    for m in RE_TEL_TXT.finditer(limpo):
        d = re.sub(r"\D", "", m.group())
        if len(d) in (12, 13) and d.startswith("55"):
            d = d[2:]
        if len(d) in (10, 11) and d not in achados:
            achados.append(d)

    if not achados:
        d = re.sub(r"\D", "", limpo)
        if len(d) in (12, 13) and d.startswith("55"):
            d = d[2:]
        if len(d) in (8, 9, 10, 11):
            achados = [d]

    # o que sobra de texto depois de tirar numeros e pontuacao
    resto = RE_TEL_TXT.sub(" ", limpo)
    resto = RE_CNPJ_TXT.sub(" ", resto)
    resto = re.sub(r"[\d()+.\-:|/;,]+", " ", resto)
    nomes = re.sub(r"\s{2,}", " ", resto).strip(" -/|:")

    notas = []
    if nomes:
        notas.append("contato citado no campo telefone: " + nomes)
    if len(achados) > 1:
        notas.append("outros telefones: " + ", ".join(achados[1:]))
    if RE_CNPJ_TXT.search(s):
        notas.append("havia um CNPJ no campo telefone: " + RE_CNPJ_TXT.search(s).group())

    principal = achados[0] if achados else ""
    if not principal and not notas and s:
        notas.append("campo telefone trazia: " + s)

    return principal, ("; ".join(notas) if notas else "")

def anexar_obs(obs, nota):
    if not nota:
        return obs
    marca = "[migração 16/09/2026] " + nota
    if not obs:
        return marca
    return obs + "\n\n" + marca

def texto(v):
    return trim(v)

def memo(v):
    if v is None:
        return None
    s = str(v).replace(" ", " ").strip()
    return s if s else None

def inteiro(v):
    s = trim(v)
    if s == "":
        return None
    try:
        return int(float(s.replace(",", ".")))
    except ValueError:
        return None

def moeda(v):
    if v is None or trim(v) == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    s = trim(v).replace("R$", "").replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return None

def data(v):
    if v is None or trim(v) == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = trim(v)
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None

def hora(v):
    if v is None or trim(v) == "":
        return None
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    return trim(v)

def vazio(v):
    return v is None or (isinstance(v, str) and v.strip() == "")

# ---------------------------------------------------------------- leitura
def cabecalhos(ws):
    """{nome normalizado da coluna: indice base 0}"""
    mapa = {}
    for i, c in enumerate(next(ws.iter_rows(min_row=LINHA_CABECALHO,
                                            max_row=LINHA_CABECALHO,
                                            values_only=True))):
        nome = trim(c)
        if nome:
            chave = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode().upper()
            chave = re.sub(r"[^A-Z0-9]+", "_", chave).strip("_")
            if chave not in mapa:
                mapa[chave] = i
    return mapa

def pega(linha, mapa, chave):
    i = mapa.get(chave)
    if i is None or i >= len(linha):
        return None
    return linha[i]

def exige(mapa, chaves, aba):
    faltam = [c for c in chaves if c not in mapa]
    if faltam:
        raise SystemExit("ERRO: colunas ausentes na aba %s: %s\nEncontradas: %s"
                         % (aba, faltam, sorted(mapa.keys())))

# ---------------------------------------------------------------- extracao
def extrair(caminho_xlsm, pasta_saida):
    wb = openpyxl.load_workbook(caminho_xlsm, read_only=True, data_only=True)
    rel = {"origem": os.path.basename(caminho_xlsm),
           "extraido_em": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
           "normalizacoes": {}, "descartes": [], "avisos": [], "excecoes": []}

    def conta_norm(campo):
        rel["normalizacoes"][campo] = rel["normalizacoes"].get(campo, 0) + 1

    # ---------------- LISTAS ----------------
    ws = wb["Listas"]
    listas = []
    linhas = list(ws.iter_rows(min_row=2, max_row=40, values_only=True))
    tipos = linhas[0]
    for col, tipo in enumerate(tipos):
        t = trim(tipo)
        if not t or t.upper().startswith("FILTRO"):
            continue
        ordem = 0
        for lin in linhas[1:]:
            if col >= len(lin):
                continue
            v = trim(lin[col])
            if v:
                ordem += 1
                listas.append({"tipo": t, "valor": v, "ordem": ordem, "ativo": True})

    # ---------------- CLIENTES ----------------
    ws = wb["Clientes"]
    m = cabecalhos(ws)
    exige(m, ["COD_CLIENTE", "EMPRESA", "CIDADE", "ESTADO", "SEGMENTO",
              "TELEFONE_GERAL", "E_MAIL_GERAL", "CNPJ", "OBSERVACOES",
              "RESPONSAVEL", "ADERENCIA", "PORTE", "QUALIFICACAO"], "Clientes")
    clientes, por_codigo = [], {}
    seq = 0
    for lin in ws.iter_rows(min_row=LINHA_DADOS, values_only=True):
        emp_bruto = pega(lin, m, "EMPRESA")
        if vazio(emp_bruto):
            continue
        seq += 1
        cod = inteiro(pega(lin, m, "COD_CLIENTE"))
        emp = maiusc(emp_bruto)
        if emp != trim(emp_bruto):
            conta_norm("clientes.empresa")
        cid_bruto = pega(lin, m, "CIDADE")
        cid = maiusc(cid_bruto)
        if cid and cid != trim(cid_bruto):
            conta_norm("clientes.cidade")
        mail_bruto = pega(lin, m, "E_MAIL_GERAL")
        mail = minusc(mail_bruto)
        if mail and mail != trim(mail_bruto):
            conta_norm("clientes.email")
        tel_bruto = pega(lin, m, "TELEFONE_GERAL")
        tel, nota_tel = telefone_partes(tel_bruto)
        if nota_tel:
            conta_norm("clientes.telefone_para_obs")
        if tel and tel != trim(tel_bruto):
            conta_norm("clientes.telefone")
        cnpj_bruto = pega(lin, m, "CNPJ")
        cnpj = digitos(cnpj_bruto)
        if cnpj and cnpj != trim(cnpj_bruto):
            conta_norm("clientes.cnpj")
        if cnpj and len(cnpj) != 14:
            rel["avisos"].append("CNPJ com %d digitos em %s" % (len(cnpj), emp))
            cnpj = ""
        ader = inteiro(pega(lin, m, "ADERENCIA"))
        if ader is None and not vazio(pega(lin, m, "ADERENCIA")):
            conta_norm("clientes.aderencia_descartada")
        porte = inteiro(pega(lin, m, "PORTE"))
        reg = {
            "seq": seq,
            "codigo_cliente": cod,
            "estagio": "Cliente" if cod else "Pré-cliente",
            "empresa": emp,
            "cnpj": cnpj or None,
            "cidade": cid or None,
            "uf": maiusc(pega(lin, m, "ESTADO"))[:2] or None,
            "segmento": texto(pega(lin, m, "SEGMENTO")) or None,
            "telefone": tel or None,
            "email": mail or None,
            "responsavel": texto(pega(lin, m, "RESPONSAVEL")) or None,
            "aderencia": ader if ader in (1, 2, 3) else None,
            "porte": porte if porte in (1, 2, 3) else None,
            "qualificacao": texto(pega(lin, m, "QUALIFICACAO")) or None,
            "observacoes": anexar_obs(memo(pega(lin, m, "OBSERVACOES")) or "", nota_tel) or None,
            "ativo": True,
        }
        clientes.append(reg)
        if cod:
            if cod in por_codigo:
                raise SystemExit("ERRO: CÓD. CLIENTE duplicado: %s" % cod)
            por_codigo[cod] = seq

    # ---------------- CONTATOS ----------------
    ws = wb["Contatos"]
    m = cabecalhos(ws)
    exige(m, ["COD_CONTATO", "COD_CLIENTE", "CONTROLE", "CONTATO",
              "CARGO", "TELEFONE", "E_MAIL", "OBSERVACOES"], "Contatos")
    contatos, por_cod_contato = [], {}
    seq = 0
    for lin in ws.iter_rows(min_row=LINHA_DADOS, values_only=True):
        codigo = trim(pega(lin, m, "COD_CONTATO"))
        if not codigo:
            continue
        cod_cli = inteiro(pega(lin, m, "COD_CLIENTE"))
        if cod_cli not in por_codigo:
            raise SystemExit("ERRO: contato %s aponta para cliente inexistente: %s" % (codigo, cod_cli))
        seq += 1
        nome_cto = texto(pega(lin, m, "CONTATO"))
        if not nome_cto:
            nome_cto = "GERAL"
            rel["excecoes"].append("contato %s sem nome -> gravado como GERAL (convencao ja usada em outros registros)" % codigo)
        mail_bruto = pega(lin, m, "E_MAIL")
        mail = minusc(mail_bruto)
        if mail and mail != trim(mail_bruto):
            conta_norm("contatos.email")
        tel_cto, nota_cto = telefone_partes(pega(lin, m, "TELEFONE"))
        if nota_cto:
            conta_norm("contatos.telefone_para_obs")
        reg = {
            "seq": seq,
            "codigo": codigo,
            "controle": inteiro(pega(lin, m, "CONTROLE")),
            "cliente_seq": por_codigo[cod_cli],
            "nome": nome_cto,
            "cargo": texto(pega(lin, m, "CARGO")) or None,
            "telefone": tel_cto or None,
            "email": mail or None,
            "observacoes": anexar_obs(memo(pega(lin, m, "OBSERVACOES")) or "", nota_cto) or None,
            "ativo": True,
        }
        contatos.append(reg)
        por_cod_contato[codigo] = seq

    # ---------------- OPORTUNIDADES ----------------
    ws = wb["Oportunidades"]
    m = cabecalhos(ws)
    exige(m, ["ATENDIMENTO", "COD_CONTATO", "COD_CLIENTE", "CONTROLE", "ETAPA",
              "ORCAMENTO", "RESPONSAVEL", "MAQUINA_DESCRICAO", "FAMILIA_DE_MAQUINA",
              "VALOR_R", "ORIGEM", "PRIORIDADE", "DATA_DE_ENTRADA", "DATA_DA_PROPOSTA",
              "ULTIMA_INTERACAO", "TENTATIVAS", "PROXIMA_ACAO", "DATA_PROX_ACAO",
              "RETOMAR_EM", "MOTIVO_DE_DESFECHO", "DATA_DO_DESFECHO", "OBSERVACOES",
              "CATEGORIA", "TIPO_DE_VENDA", "ATENDIMENTO_ANTERIOR"], "Oportunidades")
    oportunidades = []
    seq = 0
    for lin in ws.iter_rows(min_row=LINHA_DADOS, values_only=True):
        codigo = trim(pega(lin, m, "ATENDIMENTO"))
        if not codigo:
            continue
        cod_cto = trim(pega(lin, m, "COD_CONTATO"))
        if cod_cto not in por_cod_contato:
            raise SystemExit("ERRO: atendimento %s aponta para contato inexistente: %s" % (codigo, cod_cto))
        cod_cli = inteiro(pega(lin, m, "COD_CLIENTE"))
        seq += 1
        maq_bruto = pega(lin, m, "MAQUINA_DESCRICAO")
        maq = maiusc(maq_bruto)
        if maq and maq != trim(maq_bruto):
            conta_norm("oportunidades.maquina")
        etapa = texto(pega(lin, m, "ETAPA"))
        if not etapa:
            rel["excecoes"].append("atendimento %s sem etapa -> nao migrado" % codigo)
            seq -= 1
            continue
        reg = {
            "seq": seq,
            "codigo": codigo,
            "controle": inteiro(pega(lin, m, "CONTROLE")),
            "contato_seq": por_cod_contato[cod_cto],
            "cliente_seq": por_codigo[cod_cli],
            "codigo_anterior": trim(pega(lin, m, "ATENDIMENTO_ANTERIOR")) or None,
            "orcamento": maiusc(pega(lin, m, "ORCAMENTO")) or None,
            "etapa": etapa,
            "responsavel": texto(pega(lin, m, "RESPONSAVEL")) or None,
            "maquina": maq or None,
            "familia": texto(pega(lin, m, "FAMILIA_DE_MAQUINA")) or None,
            "categoria": texto(pega(lin, m, "CATEGORIA")) or None,
            "tipo_venda": texto(pega(lin, m, "TIPO_DE_VENDA")) or None,
            "valor": moeda(pega(lin, m, "VALOR_R")),
            "origem": texto(pega(lin, m, "ORIGEM")) or None,
            "prioridade": texto(pega(lin, m, "PRIORIDADE")) or None,
            "dt_entrada": data(pega(lin, m, "DATA_DE_ENTRADA")),
            "dt_proposta": data(pega(lin, m, "DATA_DA_PROPOSTA")),
            "ultima_interacao": data(pega(lin, m, "ULTIMA_INTERACAO")),
            "tentativas": inteiro(pega(lin, m, "TENTATIVAS")),
            "prox_acao": texto(pega(lin, m, "PROXIMA_ACAO")) or None,
            "dt_prox_acao": data(pega(lin, m, "DATA_PROX_ACAO")),
            "retomar_em": data(pega(lin, m, "RETOMAR_EM")),
            "motivo_desfecho": texto(pega(lin, m, "MOTIVO_DE_DESFECHO")) or None,
            "dt_desfecho": data(pega(lin, m, "DATA_DO_DESFECHO")),
            "observacoes": memo(pega(lin, m, "OBSERVACOES")),
        }
        oportunidades.append(reg)

    # resolver retomadas por codigo
    idx_por_codigo = {o["codigo"]: o["seq"] for o in oportunidades}
    for o in oportunidades:
        o["anterior_seq"] = idx_por_codigo.get(o["codigo_anterior"]) if o["codigo_anterior"] else None
        if o["codigo_anterior"] and o["anterior_seq"] is None:
            rel["avisos"].append("atendimento anterior nao encontrado: %s em %s"
                                 % (o["codigo_anterior"], o["codigo"]))

    # ---------------- limites do DDL ----------------
    # Descobrir estouro na carga custa uma migracao inteira. Confere aqui.
    LIMITES = {
        "clientes": {"estagio": 12, "empresa": 120, "cnpj": 20, "cidade": 60, "uf": 2,
                     "segmento": 60, "telefone": 30, "email": 120, "responsavel": 60,
                     "qualificacao": 40},
        "contatos": {"codigo": 15, "nome": 120, "cargo": 60, "telefone": 30, "email": 120},
        "oportunidades": {"codigo": 15, "orcamento": 30, "etapa": 40, "responsavel": 60,
                          "maquina": 255, "familia": 60, "categoria": 60, "tipo_venda": 40,
                          "origem": 40, "prioridade": 20, "motivo_desfecho": 60},
    }
    estouros = []
    for tabela, regs in (("clientes", clientes), ("contatos", contatos), ("oportunidades", oportunidades)):
        for reg in regs:
            for campo, limite in LIMITES[tabela].items():
                v = reg.get(campo)
                if v is not None and len(str(v)) > limite:
                    estouros.append("%s.%s: %d caracteres (limite %d) em %s"
                                    % (tabela, campo, len(str(v)), limite,
                                       reg.get("codigo") or reg.get("empresa")))
    if estouros:
        print("ERRO: valor maior que a coluna do banco. Corrija o DDL ou o dado:")
        for e in estouros[:20]:
            print("  -", e)
        raise SystemExit(1)

    # ---------------- esperado (conferencia) ----------------
    etapas = {}
    for o in oportunidades:
        etapas[o["etapa"]] = etapas.get(o["etapa"], 0) + 1
    esperado = {
        "clientes": len(clientes),
        "clientes_com_codigo": sum(1 for c in clientes if c["codigo_cliente"]),
        "clientes_pre": sum(1 for c in clientes if not c["codigo_cliente"]),
        "contatos": len(contatos),
        "oportunidades": len(oportunidades),
        "listas": len(listas),
        "soma_valor": round(sum(o["valor"] or 0 for o in oportunidades), 2),
        "max_codigo_cliente": max((c["codigo_cliente"] or 0) for c in clientes),
        "max_controle_contato": max((c["controle"] or 0) for c in contatos),
        "max_controle_oportunidade": max((o["controle"] or 0) for o in oportunidades),
        "por_etapa": dict(sorted(etapas.items(), key=lambda x: -x[1])),
    }

    os.makedirs(pasta_saida, exist_ok=True)
    def grava(nome, obj):
        caminho = os.path.join(pasta_saida, nome)
        with open(caminho, "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, indent=1)
        return "%-24s %6d registros  %8.1f KB" % (nome, len(obj) if isinstance(obj, list) else 1,
                                                  os.path.getsize(caminho) / 1024)

    print(grava("carga-listas.json", listas))
    print(grava("carga-clientes.json", clientes))
    print(grava("carga-contatos.json", contatos))
    print(grava("carga-oportunidades.json", oportunidades))
    print(grava("esperado.json", esperado))
    rel["esperado"] = esperado
    print(grava("relatorio-extracao.json", rel))

    print()
    print("ESPERADO:")
    for k, v in esperado.items():
        if k != "por_etapa":
            print("  %-28s %s" % (k, v))
    print("  por_etapa:")
    for k, v in esperado["por_etapa"].items():
        print("     %-26s %d" % (k, v))
    print()
    print("NORMALIZACOES APLICADAS:")
    for k, v in sorted(rel["normalizacoes"].items()):
        print("  %-34s %d registro(s)" % (k, v))
    if rel["excecoes"]:
        print()
        print("EXCECOES (%d) - conferir antes da virada:" % len(rel["excecoes"]))
        for e in rel["excecoes"][:30]:
            print("  -", e)
    if rel["avisos"]:
        print()
        print("AVISOS (%d):" % len(rel["avisos"]))
        for a in rel["avisos"][:20]:
            print("  -", a)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit("uso: python3 extrair.py <arquivo.xlsm> <pasta de saida>")
    extrair(sys.argv[1], sys.argv[2])
